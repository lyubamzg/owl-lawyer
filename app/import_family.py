#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Раздел «Семейное право»: сканы документов → анкета бракоразводного дела.

Складывает сканы в папку «scans-family» рядом со скриптом:
    паспорт истца.pdf          — паспорт (распознаётся только ИСТЦА)
    свидетельство о браке.pdf  — свидетельство о заключении брака
    рождение-1.pdf ... рождение-3.pdf — свидетельства о рождении детей

OCR распознаёт текст, парсеры вычленяют поля и заполняют «family-questionnaire.xlsx».
Правило заполнения: в анкету пишется только то, что ПУСТО — ручные правки юриста
не трогаются. Окончательная сверка по оригиналу всегда за человеком.

    python3 import_family.py            — загрузка сканов из «scans-family»
    python3 import_family.py файл1 ...  — конкретные файлы
    python3 import_family.py --демо      — создать анкету с тестовой строкой
    python3 import_family.py --демо-парсер — проверить парсеры на тестовых текстах

Данные никуда не отправляются — всё распознаётся на этом компьютере.
"""
import re
import subprocess
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

import platform_utils

BASE = Path(__file__).resolve().parent
ANKETA = BASE / 'family-questionnaire.xlsx'
SRC_DIR = BASE / 'scans-family'
OUT_DIR = BASE / 'output' / 'ocr-family'

IMAGE_EXT = {'.jpg', '.jpeg', '.png', '.tif', '.tiff', '.bmp', '.pdf'}
# электронные документы: Word (doc/docx/rtf/odt), Excel (xlsx/xls/ods), текст
DOC_EXT = IMAGE_EXT | {'.docx', '.doc', '.rtf', '.odt',
                       '.xlsx', '.xls', '.ods', '.txt', '.csv'}

# ------------------------------------------------------------------ regex ----
# База та же, что в ocr_debtors.py / import_housing.py, но с учётом капризов
# OCR на бланках ЗАГС (римские серии распознаются латиницей, «МЮ» как «MIO»).
FIO_RE = re.compile(
    r'([А-ЯЁ][а-яё\-]+ +[А-ЯЁ][а-яё\-]+ +[А-ЯЁ][а-яё\-]+(?:ович|евич|овна|евна|ична|инична))\b')
# «широкий» вариант: 3 любых слова с заглавной (капс OCR, родительный падеж в бланках ЗАГС,
# строки-разрывы). Разделитель \s+ — бланки ломают строки где угодно.
FIO_ANY_RE = re.compile(r'\b([А-ЯЁ][А-ЯЁа-яё\-]{2,}\s+[А-ЯЁ][А-ЯЁа-яё\-]{2,}\s+[А-ЯЁ][А-ЯЁа-яё\-]{2,})\b')
DATE_RE = re.compile(r'\b(\d{2}[.,]\d{2}[.,](19|20)\d{2})\b')
PASPORT_RE = re.compile(r'\b(\d{2} ?\d{2})[\s№]*(\d{6})\b')
KEM_VIDAN_RE = re.compile(r'выдан[^\n]{10,200}', re.IGNORECASE)
ADDR_RE = re.compile(r'[Зз]арегистриров(?:ан|а)[а-яё]*\s*(?:по адресу|г\.|с\.|в\b)?[^\n]{0,20}'
                     r'([А-ЯЁ].{9,180})')
# серия свидетельства: «VII-МЮ 071684» / «VП-МЮ 071684» / «VII-MIO 071684»
SV_SERIE_RE = re.compile(r'\b([IVXLCМЮП]{1,7})\s*[-–—]\s*(МЮ|MIO|МЮ|АЮ)\s*(?:№|N|№|\bn\b)?\s*(\d{6,7})\b')
ACT_RE = re.compile(r'запись акта\s*(?:о заключении брака|о рождении)?[^\n]{0,40}?'
                    r'(?:№|N)\s*(\d[\d/]*)', re.IGNORECASE)
OT_RE = re.compile(r'(?:№|N)\s*\d[\d/]*\s*[^0-9]{0,30}?от\s+(\d{2}[.,]\d{2}[.,](19|20)\d{2})', re.IGNORECASE)
ZAGS_RE = re.compile(r'[^\n]{0,10}(?:отдел|управлен|дворец|центр)[^\n]{0,140}ЗАГС[^\n]{0,60}',
                     re.IGNORECASE)

LAT_TO_CYR = str.maketrans({'V': 'V', 'П': 'П'})  # серии нормализует fix_serie


def fix_serie(txt):
    """Нормализует испорченные OCR серии: латиница → кириллица, МIO→МЮ."""
    t = txt.replace('MIO', 'МЮ').replace('MЮ', 'МЮ').replace('MI0', 'МЮ')
    t = t.replace('-', '–').replace('—', '–')
    return t


def norm_date(s):
    return re.sub(r'(\d{2})[.,](\d{2})[.,](\d{4})', r'\1.\2.\3', s)


# ------------------------------------------------------------------- OCR ----
def run_tesseract(img: Path, out_dir: Path = None) -> str:
    out_dir = Path(out_dir) if out_dir else img.parent
    out = out_dir / platform_utils.ascii_filename(img.stem)
    subprocess.run(['tesseract', str(img), str(out), '-l', 'rus', '--psm', '4'],
                   check=True, capture_output=True, env=platform_utils.child_env())
    return out.with_suffix('.txt').read_text(encoding='utf-8')


def ocr_doc(path: Path, out_dir: Path = None) -> str:
    """Распознаёт скан (pdf → страницы pdftoppm, картинка → напрямую), tesseract."""
    out_dir = Path(out_dir) if out_dir else OUT_DIR
    out_dir.mkdir(parents=True, exist_ok=True)
    texts = []
    if path.suffix.lower() == '.pdf':
        stem = platform_utils.ascii_filename(path.stem)
        subprocess.run(['pdftoppm', '-r', '300', '-gray', '-png',
                        str(path), str(out_dir / stem)], check=True,
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        for img in sorted(out_dir.glob(f'{stem}-*-*.png') or out_dir.glob(f'{stem}-*.png')):
            texts.append(run_tesseract(img))
    else:
        texts.append(run_tesseract(path, out_dir))
    return '\n'.join(texts)


# ------------------------------------------------------- офисные документы ----
OFFICE_CONV_EXT = {'.rtf', '.odt', '.ods', '.xls', '.doc'}


def soffice_convert(src_path: Path, out_ext: str, out_dir: Path) -> Path:
    """Конвертирует офиц. документ (rtf/odt/ods/doc) в txt или xlsx — LibreOffice."""
    out_dir.mkdir(parents=True, exist_ok=True)
    subprocess.run([platform_utils.soffice_cmd(), '--headless', '--convert-to', out_ext,
                    '--outdir', str(out_dir), str(src_path)],
                   check=True, capture_output=True, env=platform_utils.child_env())
    produced = out_dir / (src_path.stem + '.' + out_ext)
    if not produced.exists():  # LibreOffice меняет расширение по своему решению
        found = sorted(out_dir.glob(src_path.stem + '.*'))
        produced = found[-1] if found else produced
    if produced.exists():
        destination = produced.with_name(platform_utils.ascii_filename(produced.name))
        if destination != produced:
            produced.rename(destination)
            produced = destination
    return produced


def xlsx_to_text(path: Path) -> str:
    """Все листы книги openpyxl → текст «Лист: ячейки» (для журнала и парсеров)."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    chunks = []
    for wsname in wb.sheetnames:
        ws = wb[wsname]
        chunks.append(f'=== Лист: {wsname} ===')
        for row in ws.iter_rows(values_only=True):
            vals = [str(v) for v in row if v not in (None, '')]
            if vals:
                chunks.append(' | '.join(vals))
    wb.close()
    return '\n'.join(chunks)


def docx_to_text(path: Path) -> str:
    """Абзацы и ячейки таблиц .docx (python-docx)."""
    from docx import Document
    doc = Document(str(path))
    chunks = [p.text for p in doc.paragraphs if p.text.strip()]
    for tbl in doc.tables:
        for row in tbl.rows:
            vals = [c.text.strip() for c in row.cells if c.text.strip()]
            if vals:
                chunks.append(' | '.join(vals))
    return '\n'.join(chunks)


def doc_text(path: Path, out_dir: Path = None) -> str:
    """Текст ЛЮБОГО поддержанного документа: сканы — OCR, офисные — напрямую.

    pdf/картинки → tesseract (как раньше); docx/doc/rtf/odt → Word-текст;
    xlsx/xls/ods → таблицы; txt/csv → как есть. Для офисных OCR не нужен.
    """
    path = Path(path)
    out_dir = Path(out_dir) if out_dir else OUT_DIR
    ext = path.suffix.lower()
    if ext == '.txt' or ext == '.csv':
        return path.read_text(encoding='utf-8', errors='replace')
    if ext == '.docx':
        return docx_to_text(path)
    if ext == '.xlsx':
        return xlsx_to_text(path)
    if ext == '.doc':
        try:
            out = subprocess.run(['antiword', str(path)], check=True,
                                 capture_output=True, text=True,
                                 env=platform_utils.child_env())
            return out.stdout
        except Exception:  # antiword нет/не взял — отдаём в LibreOffice
            out_path = soffice_convert(path, 'txt', out_dir)
            return out_path.read_text(encoding='utf-8', errors='replace')
    if ext in OFFICE_CONV_EXT:
        conv = soffice_convert(path, 'txt' if ext in ('.rtf', '.odt') else 'xlsx',
                               out_dir)
        if conv.suffix == '.xlsx':
            return xlsx_to_text(conv)
        return conv.read_text(encoding='utf-8', errors='replace')
    return ocr_doc(path, out_dir)  # сканы: pdf и картинки — по прежнему конвейеру


# ------------------------------------------------------------ классификация ----
def classify_doc(text, filename):
    """Тип документа: по имени файла, затем по шапкам текста."""
    f = filename.lower()
    if 'пасп' in f:
        return 'паспорт'
    if 'брак' in f or 'жени' in f:
        return 'брак'
    if 'рожд' in f or 'ребён' in f or 'ребен' in f:
        return 'рождение'
    t = text.upper()
    if 'О ЗАКЛЮЧЕНИИ БРАКА' in t or 'БРАКА' in t and 'СВИДЕТЕЛЬСТВО' in t:
        return 'брак'
    if 'О РОЖДЕНИИ' in t or 'РОЖДЕНИИ' in t and 'СВИДЕТЕЛЬСТВО' in t:
        return 'рождение'
    if 'ПАСПОРТ' in t or 'КЕМ ВЫДАН' in t:
        return 'паспорт'
    return ''


# ---------------------------------------------------------------- парсеры ----
def first_person_date(fio_match, text):
    """Дата рождения: ближайшая после ФИО, не дата выдачи паспорта."""
    after = text[fio_match.end():fio_match.end() + 80]
    for dm in DATE_RE.finditer(after):
        tail = after[dm.start():dm.start() + 40]
        if KEM_VIDAN_RE.search(tail) or 'выдан' in tail[:12]:
            continue
        return dm.group(1).replace(',', '.')
    return ''


def fix_caps(w):
    """Капс OCR «ИВАНОВА» → «Иванова»; смешанные варианты не трогаем."""
    return (w.capitalize() if len(w) > 2 and w.isupper() else w)


def fix_fio(s):
    return ' '.join(fix_caps(w) for w in re.sub(r'\s+', ' ', s).strip().split())


def passport_fio(text):
    """ФИО из паспорта: бланк печатает Фамилия/Имя/Отчество с новой строки капсом."""
    norm = re.sub(r'(Фамилия|Имя|Отчество)\s*\n?\s*', ' ', text)
    m = FIO_ANY_RE.search(norm)
    return fix_fio(m.group(1)) if m else ''


def parse_passport(text):
    """→ {fio, dr, pasport, kem_kogda, addr} для паспорта ИСТЦА."""
    out = {'fio': '', 'dr': '', 'pasport': '', 'kem_kogda': '', 'addr': ''}
    norm = re.sub(r'(Фамилия|Имя|Отчество)\s*\n?\s*', ' ', text)
    out['fio'] = passport_fio(text)
    m = FIO_ANY_RE.search(norm)
    if m:
        out['dr'] = first_person_date(m, norm)
        # паспорт ищем до и после ФИО (как в ocr_debtors.extract_persons)
        zone = norm[max(0, m.start() - 100):m.end() + 160]
        p = PASPORT_RE.search(zone)
        if p:
            out['pasport'] = f'{p.group(1)} {p.group(2)}'
    km = KEM_VIDAN_RE.search(text)
    if km:
        out['kem_kogda'] = ' '.join(km.group(0).split()).rstrip('.,;')
    am = ADDR_RE.search(text) or re.search(r'(?i)[зЗ]арегистриров.*?[:—-]?\s*(.{9,180})', text)
    if am:
        out['addr'] = ' '.join(am.group(1).split()).strip(' .,;')
    return out


def parse_sv_obsh(text):
    """Общее для свидетельств: серия, номер записи акта, дата, ЗАГС."""
    out = {'seria': '', 'act_num': '', 'act_date': '', 'zags': ''}
    sm = SV_SERIE_RE.search(fix_serie(text))
    if sm:
        out['seria'] = f'{sm.group(1)}-МЮ {sm.group(3)}'
    tm = ACT_RE.search(text)
    if tm:
        out['act_num'] = tm.group(1)
        om = OT_RE.search(text[max(0, tm.start() - 10):tm.end() + 120])
        if om:
            out['act_date'] = om.group(1).replace(',', '.')
    if not out['act_date']:
        dm = DATE_RE.search(text)
        if dm:
            out['act_date'] = dm.group(1).replace(',', '.')
    zm = ZAGS_RE.search(text)
    if zm:
        out['zags'] = ' '.join(zm.group(0).split()).strip(' .,;')
    return out


def parse_sv_braka(text):
    """→ серия, номер записи, дата брака, ЗАГС; отдельно — оба супруга (для сверки)."""
    out = parse_sv_obsh(text)
    out['suprugi'] = ', '.join(m.group(1) for m in list(FIO_RE.finditer(text))[:2])
    if not out['act_date']:
        # дата рождения одного из супругов — не она; ищем дату у слов «брак»
        bm = re.search(r'(?i)брак[^\n]{0,200}?(\d{2}[.,]\d{2}[.,](?:19|20)\d{2})', text)
        if bm:
            out['act_date'] = bm.group(1).replace(',', '.')
    return out


FIO_STOP = {'отдел', 'управление', 'дворец', 'город', 'район', 'комитет', 'загс'}


def fio_candidates(text):
    """Кандидаты ФИО «широким» regex без слов-ловушек (Отдел/Управление/ЗАГС…)."""
    out = []
    for m in FIO_ANY_RE.finditer(text):
        if m.group(1).split()[0].lower() in FIO_STOP:
            continue
        out.append(m)
    return out


def parse_sv_rozhd(text):
    """→ список словарей {fio, dr, seria, act_num, act_date, zags} — по одному на бланк.

    В одном файле может быть несколько свидетельств: каждая веха «Родивш…» — ребёнок.
    """
    rozh = []
    heads = list(re.finditer(r'(?i)Родивш[a-яё]+\s+', text))
    for i, hd in enumerate(heads):
        seg = text[hd.start(): heads[i + 1].start() if i + 1 < len(heads) else len(text)]
        after = text[hd.end():hd.end() + 180]
        wm = FIO_ANY_RE.search(after)
        if not wm:
            continue
        base = parse_sv_obsh(seg)
        d = {'fio': fix_fio(wm.group(1)), 'dr': first_person_date(wm, after)}
        d.update(base)
        rozh.append(d)
    if not rozh:  # нет вехи «Родивш…» — один ребёнок по общим правилам
        cands = fio_candidates(text)
        if cands:
            m = cands[0]
            d = {'fio': fix_fio(m.group(1)), 'dr': first_person_date(m, text)}
            d.update(parse_sv_obsh(text))
            rozh.append(d)
    return rozh


# ---------------------------------------------------------------- анкета ----
thin = Side(style='thin', color='999999')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FILL = PatternFill('solid', fgColor='DDEBF7')
TEST_FILL = PatternFill('solid', fgColor='FFF2CC')
HEAD_FONT = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical='top')

DELO_COLS = ['№', 'Суд (шапка)', 'Госпошлина, руб.', 'Примечание',
             'Истец ФИО', 'Истец дата рождения', 'Истец адрес регистрации', 'Истец телефон',
             'Паспорт истца (кем и когда выдан)',
             'Свидетельство о браке: серия', 'Номер записи акта о браке',
             'Дата заключения брака', 'Наименование ЗАГС',
             'Дата прекращения совместной жизни',
             'Ответчик ФИО', 'Дата рождения ответчика', 'Адрес ответчика',
             'Паспорт ответчика (кем и когда выдан)',
             'Доля алиментов (1/4, 1/3, 1/2 или твердая сумма)',
             'Дети проживают с (истец/ответчик)']
DETI_COLS = ['Номер ребёнка (1-3)', 'ФИО ребёнка', 'Дата рождения',
             'Свид-во о рождении (серия, номер записи)', 'Примечание']
SKAN_COLS = ['Файл', 'Тип документа', 'Что извлечено', 'Проверено (✓)']

DEMO = {
    'Дело': ['1', 'В Примерный городской суд\nПримерного района Примерной области', '',
             'ТЕСТОВАЯ СТРОКА — удалите и вносите реальное дело',
             'Иванова Анна Петровна', '14.02.1986', 'Примерная обл., г. Пример, ул. Лесная, д. 5, кв. 12',
             '8-911-000-11-22', '4010 123456 МВД России по Примерной области от 12.03.2008 г.',
             'III-МЮ 077777', '777', '20.06.2010', 'Отдел ЗАГС Примерного района Примерной области',
             '01.10.2025', 'Иванов Ирик Игоревич', '02.07.1984',
             'Примерная обл., г. Пример, ул. Лесная, д. 5, кв. 12',
             '4105 775533 ТП № 34 отд. МВД по Примерной области от 24.09.2004 г.',
             '1/3', 'истец'],
    'Дети': [[1, 'Иванов Никита Ирикович', '03.05.2012', 'VII-MЮ 095123, № 123', ''],
             [2, 'Иванова Полина Ириковна', '11.11.2015', 'VIII-МЮ 098456, № 456', '']],
}
DEMO_TEXTS = {
    'паспорт': 'ПАСПОРТ РФ\nФамилия ИВАНОВА\nИмя Анна\nОтчество Петровна\n'
               "14.02.1986 г.р., м.р. г. Колпино\n4010 123456 выдан 12.03.2008 г. "
               'МВД России по Примерной области\n'
               'Зарегистрирована по адресу: Примерная обл., г. Пример, ул. Лесная, д. 5, кв. 12',
    'брак': 'СВИДЕТЕЛЬСТВО О ЗАКЛЮЧЕНИИ БРАКА\nVII-MIO 071684\n'
            'Отдел ЗАГС Примерного района Примерной области\n'
            'запись акта о заключении брака № 258 от 20.06.2010\n'
            'Иванов Ирик Игоревич и Иванова Анна Петровна заключили брак.',
    'рождение': 'СВИДЕТЕЛЬСТВО О РОЖДЕНИИ\nVIII-МЮ 098456\n'
                'Отдел ЗАГС Примерного района\nзапись акта о рождении № 456 от 20.11.2015\n'
                'Родившейся Ивановой Полине Ириковне, 11.11.2015 г.р.',
}


def ensure_anketa(path: Path = ANKETA, demo: bool = False):
    """Создаёт анкету, если её нет. Существующую не трогает."""
    if path.exists():
        return False
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Инструкция'
    lines = [
        ['АНКЕТА БРАКОРАЗВОДНОГО ДЕЛА — ИНСТРУКЦИЯ'],
        [''],
        ['Как работать:'],
        ['1. Сканы (паспорт ИСТЦА, свидетельство о браке, свидетельства о рождении детей)'],
        ['   кладите в папку «scans-family» рядом с приложением:'],
        ['   «паспорт истца», «свидетельство о браке», «рождение-1», «рождение-2»…'],
        ['2. В приложении — раздел «Семейное право» → «Загрузить сканы документов».'],
        ['   OCR распознает и впишет данные ТУЛЬКО в пустые ячейки — ваши правки не сотрёт.'],
        ['3. Дозаполните лист «Дело» руками: суд, данные ответчика, доля алиментов,'],
        ['   дата прекращения совместной жизни. Сверьте с оригиналом всё, что вписал OCR'],
        ['   (лист «Сканы» показывает, что извлечено из каждого скана).'],
        ['4. Детей вписывайте на листе «Дети» (строка на ребёнка, номер 1–3).'],
        ['5. В приложении — «Собрать исковое заявление». Проверьте «validation-report-family.txt».'],
        [''],
        ['Форматы: даты ДД.ММ.ГГГГ; доля алиментов «1/4», «1/3», «1/2» или сумма;'],
        ['«Суд (шапка)» — как в шапке иска, две строки через Alt+Enter.'],
        ['Госпошлину программа подсказывает 5 000 ₽ (пп. 5 п. 1 ст. 333.19 НК РФ),'],
        ['истец по требованиям об алиментах освобождён (ст. 333.36 НК РФ) — сверьтесь.'],
        [''],
        ['Жёлтые строки — ТЕСТОВЫЕ ВЫМЫШЛЕННЫЕ данные. Удалите их и вносите реальные.'],
    ]
    for row in lines:
        ws.append(row)
    ws.column_dimensions['A'].width = 110
    ws['A1'].font = Font(bold=True, size=14)

    ws = wb.create_sheet('Дело')
    ws.append(DELO_COLS)
    if demo:
        ws.append(DEMO['Дело'])
    for i, w in enumerate((5, 30, 13, 24, 24, 13, 30, 16, 34, 14, 14, 13, 24, 14, 24, 13, 30, 34, 16, 14), 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.font, cell.fill, cell.alignment, cell.border = HEAD_FONT, HEAD_FILL, WRAP, BORDER
    dv = DataValidation(type='list', formula1='"истец,ответчик"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add('T2:T60')

    ws = wb.create_sheet('Дети')
    ws.append(DETI_COLS)
    if demo:
        for row in DEMO['Дети']:
            ws.append(row)
    for i, w in enumerate((14, 30, 13, 30, 24), 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.font, cell.fill, cell.alignment, cell.border = HEAD_FONT, HEAD_FILL, WRAP, BORDER
    dv2 = DataValidation(type='whole', operator='between', formula1='1', formula2='3', allow_blank=True)
    ws.add_data_validation(dv2)
    dv2.add('A2:A60')

    ws = wb.create_sheet('Сканы')
    ws.append(SKAN_COLS)
    for i, w in enumerate((36, 14, 60, 12), 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.font, cell.fill, cell.alignment, cell.border = HEAD_FONT, HEAD_FILL, WRAP, BORDER

    if demo:  # пометим тестовые строки жёлтым
        for sheet, last in (('Дело', 2), ('Дети', 3)):
            ws2 = wb[sheet]
            for row in ws2.iter_rows(min_row=2, max_row=last):
                for c in row:
                    if c.value not in (None, ''):
                        c.fill = TEST_FILL
    wb.save(path)
    return True


def cell_empty(ws, row, col):
    v = ws.cell(row=row, column=col).value
    return v is None or str(v).strip() == ''


def fill_anketa(path: Path, results):
    """Пишет распознанные данные ТОЛЬКО в пустые ячейки. Возвращает список правок."""
    wb = openpyxl.load_workbook(path)
    ws, wd, wj = wb['Дело'], wb['Дети'], wb['Сканы']
    heads = {str(c.value).strip(): c.column for c in ws[1] if c.value}
    dheati = {str(c.value).strip(): c.column for c in wd[1] if c.value}
    changed, journal = [], []

    def put_de(col_name, val):
        val = (val or '').strip()
        if not val:
            return
        col = heads.get(col_name)
        if col and cell_empty(ws, 2, col):
            ws.cell(row=2, column=col, value=val)
            changed.append(f'{col_name} ← {val}')

    doc = results.get('паспорт')
    if doc:
        fio = doc.get('fio', '')
        put_de('Истец ФИО', fio)
        put_de('Истец дата рождения', doc.get('dr', ''))
        put_de('Истец адрес регистрации', doc.get('addr', ''))
        put_de('Паспорт истца (кем и когда выдан)',
               f'{doc.get("pasport", "")} {doc.get("kem_kogda", "")}'.strip(" ,."))

    br = results.get('брак')
    if br:
        put_de('Свидетельство о браке: серия', br.get('seria', ''))
        put_de('Номер записи акта о браке', br.get('act_num', ''))
        put_de('Дата заключения брака', br.get('act_date', ''))
        put_de('Наименование ЗАГС', br.get('zags', ''))
        journal.append('брак: супруги в актовой записи — ' + br.get('suprugi', ''))

    rozh = results.get('рождение', [])
    for kid in rozh:
        num = None
        for src in results.get('_rozh_num', {}).get(kid.get('fio', ''), []):
            num = src
            break
        col_num = dheati.get('Номер ребёнка (1-3)')
        col_fio = dheati.get('ФИО ребёнка')
        target = None
        for r in range(2, 62):
            v = wd.cell(row=r, column=col_num or 1).value
            if num and str(v).strip() == str(num):
                target = r
                break
            if target is None and bool(v in (None, '')):
                target = r  # первая пустая строка
                if num:
                    wd.cell(row=target, column=col_num or 1, value=num)
                break
        if target is None:
            journal.append(f'рождение: лист «Дети» заполнен (3 ребёнка) — пропущено: {kid.get("fio", "")}')
            continue
        for col_name, val in (('ФИО ребёнка', kid.get('fio', '')),
                              ('Дата рождения', kid.get('dr', '')),
                              ('Свид-во о рождении (серия, номер записи)',
                               f'{kid.get("seria", "")}{", № " + kid["act_num"] if kid.get("act_num") else ""}')):
            val = str(val or '').strip(' ,.')
            if not val:
                continue
            col = dheati.get(col_name)
            if col and cell_empty(wd, target, col):
                wd.cell(row=target, column=col, value=val)
                changed.append(f'Дети стр. {target} · {col_name} ← {val}')

    # журнал сканов
    existing = {(str(wj.cell(row=r, column=1).value or '').strip(),
                 str(wj.cell(row=r, column=3).value or '').strip())
                for r in range(2, wj.max_row + 1)}
    for res in results.get('_rows', []):
        fname, tip, info = res
        if (fname, info) not in existing:
            r = wj.max_row + 1
            wj.cell(row=r, column=1, value=fname)
            wj.cell(row=r, column=2, value=tip)
            wj.cell(row=r, column=3, value=info)
            existing.add((fname, info))
    wb.save(path)
    return changed


def scan_files(files):
    """OCR + парсинг набора файлов. → results для fill_anketa."""
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    results = {'_rows': [], '_rozh_num': {}}
    counts = {'паспорт': 0, 'брак': 0, 'рождение': 0, '': 0}
    rozh = []
    total = len(files)
    print(f'[0/{total}] начало обработки', flush=True)
    for i, f in enumerate(files, 1):
        f = Path(f)
        try:
            text = doc_text(f, OUT_DIR)
        except Exception as ex:  # noqa: BLE001
            results['_rows'].append((f.name, '?', f'! не удалось прочитать: {ex}'))
            continue
        (OUT_DIR / platform_utils.ascii_filename(f.stem + '.txt')).write_text(text, encoding='utf-8')
        print(f'[{i}/{total}] {f.name}: готово', flush=True)
        tip = classify_doc(text, f.name)
        if tip == 'паспорт':
            p = parse_passport(text)
            results['паспорт'] = p
            info = f'ФИО: {p["fio"] or "?"}; {p["dr"] or "?"}; паспорт {p["pasport"] or "?"}; адрес: {p["addr"] or "?"}'
            counts['паспорт'] += 1
        elif tip == 'брак':
            b = parse_sv_braka(text)
            results['брак'] = b
            info = f'серия: {b["seria"] or "?"}; запись № {b["act_num"] or "?"} от {b["act_date"] or "?"}; ЗАГС: {b["zags"] or "?"}'
            counts['брак'] += 1
        elif tip == 'рождение':
            kids = parse_sv_rozhd(text)
            for i, k in enumerate(kids, 1):
                nm = re.search(r'рожд[е-]?[её]?\s*(?:-|—|\d)\s*', f.stem, re.IGNORECASE)
                m = re.search(r'(\d)', f.stem[::-1])
                num = int(m.group(1)) if m else (counts['рождение'] + i)
                results['_rozh_num'].setdefault(k.get('fio', ''), []).append(num)
            rozh += kids
            counts['рождение'] += len(kids)
            info = '; '.join(f'ребёнок: {k.get("fio") or "?"} ({k.get("dr") or "?"}, {k.get("seria") or "?"})'
                             for k in kids) or '! поля не распознаны'
            if any((k.get('fio') or '').split()[0:1] and (k['fio'].split()[0].lower().endswith(('ой', 'ей', 'е')))
                   for k in kids):
                info += ' [ФИО на бланке в родительном падеже — поправьте в анкете]'
        else:
            info = '! тип файла не опознан — переименуйте: «паспорт…», «свидетельство о браке», «рождение-1»'
        results['_rows'].append((f.name, tip or '?', info))
    if rozh:
        results['рождение'] = rozh
    return results, counts


def demo_parser():
    """Проверка парсеров на тестовых текстах (включая капризы OCR)."""
    bad = 0
    checks = [
        ('паспорт', DEMO_TEXTS['паспорт'], parse_passport(DEMO_TEXTS['паспорт']),
         {'fio': 'Иванова Анна Петровна', 'dr': '14.02.1986', 'pasport': '4010 123456'}),
        ('брак', DEMO_TEXTS['брак'], parse_sv_braka(DEMO_TEXTS['брак']),
         {'seria': 'VII-МЮ 071684', 'act_num': '258', 'act_date': '20.06.2010'}),
        ('рождение', DEMO_TEXTS['рождение'], parse_sv_rozhd(DEMO_TEXTS['рождение']),
         {'fio': 'Ивановой Полине Ириковне', 'dr': '11.11.2015'}),
        # ФИО ребёнка останется в родительном падеже, как на бланке («Родившейся…») —
        # юрист правит в анкете; программа предупредит в отчёте сборки иска.
    ]
    for tip, text, got, want in checks:
        print(f'── {tip} ──')
        if isinstance(got, list):  # рождения — список; сверяем первый бланк
            got = got[0] if got else {}
        for k, v in want.items():
            ok = got.get(k) == v
            bad += not ok
            print(f'  {"✓" if ok else "✗"} {k}: {got.get(k, "")!r} (ожидалось {v!r})'
                  if not ok else f'  ✓ {k}: {got.get(k)!r}')
    extra = {'zags': parse_sv_braka(DEMO_TEXTS['брак']).get('zags', '')}
    print(f'  ЗАГС: {extra["zags"]!r}' + (f'  ✓' if 'ЗАГС' in extra['zags'] else ' ✗ (нет «ЗАГС»)'))
    return bad


def main():
    platform_utils.fix_console()
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    flags = {a for a in sys.argv[1:] if a.startswith('--')}
    demo = '--демо' in flags or '--demo' in flags
    if '--демо-парсер' in flags:
        bad = demo_parser()
        sys.exit(0 if not bad else f'Парсеры: {bad} расхождений')

    created = ensure_anketa(demo=demo)
    if demo:
        print(f'Анкета: {ANKETA}' + (' (создана с тестовой строкой)' if created else ' — УЖЕ СУЩЕСТВОВАЛА, тестовая строка не добавлялась'))
        return
    if created:
        print('Создан анкета: «family-questionnaire.xlsx» (без тестовых строк).')

    if args:
        files = []
        for p in args:
            p = Path(p) if Path(p).exists() else BASE / p
            files += (sorted(p.iterdir()) if p.is_dir() else [p])
        files = [p for p in files if p.exists()
                 and p.suffix.lower() in DOC_EXT and not p.name.startswith('~')]
    else:
        SRC_DIR.mkdir(exist_ok=True)
        files = sorted(p for p in SRC_DIR.iterdir()
                       if p.suffix.lower() in DOC_EXT and not p.name.startswith('~'))
    if not files:
        print('Документы (jpg/png, pdf, Word: doc/docx/rtf/odt, Excel: xlsx/xls/ods, '
              'txt/csv) не найдены в выбранной папке — проверьте и повторите.')
        print(f'Положите их в папку «{SRC_DIR.name}» рядом с приложением:')
        print('  паспорт истца · свидетельство о браке · рождение-1, рождение-2 …')
        print(f'Анкета: {ANKETA}' + (' (указатели колонок готовы)' if created else ''))
        return

    print(f'РАСПОЗНАВАНИЕ СЕМЕЙНЫХ ДОКУМЕНТОВ — {len(files)} файл(ов)')
    results, counts = scan_files(files)
    for fname, tip, info in results['_rows']:
        print(f'  [{tip or "?":>8}] {fname}\n              {info}')
    changed = fill_anketa(ANKETA, results)
    print(f'Итог: паспорт {counts["паспорт"]}, брак {counts["брак"]}, дети {counts["рождение"]}')
    if not changed:
        print('Все распознанные поля уже были заполнены — анкета не изменена.')
    else:
        print('Вписано в анкету (только в ПУСТЫЕ ячейки):')
        for line in changed:
            print('  • ' + line)
    print(f'Полные тексты OCR: {OUT_DIR}')
    print('Сверьте каждую распознанную строку с оригиналом — OCR ошибается!')


if __name__ == '__main__':
    main()
