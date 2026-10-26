#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Раздел «Защита прав потребителя»: сканы → анкета потребительского дела.

Складывает сканы в папку «scans-consumer» рядом со скриптом:
    паспорт истца.pdf   — паспорт (распознаётся только ИСТЦА)
    претензия*.pdf      — досудебная претензия (текст сохраняется на лист «Сканы»)
    чек*.pdf            — платежный документ (текст — на лист «Сканы»)

OCR распознаёт паспорт истца и вписывает его данные в «consumer-questionnaire.xlsx»
(только пустые ячейки). Тексты остальных сканов сохраняются на лист «Сканы» —
суммы и даты юрист сверяет и вносит сам. Описательная часть иска (проблематика)
заполняется на листе «Проблематика» — свободный текст, каждая строка = абзац.

    python3 import_consumer.py           — загрузка сканов
    python3 import_consumer.py файлы...  — конкретные файлы
    python3 import_consumer.py --демо    — создать анкету с тестовой строкой

Данные никуда не отправляются — всё распознаётся на этом компьютере.
"""
import re
import subprocess
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

import platform_utils
import import_family as fam   # переиспользуем OCR-движок и парсер паспорта

BASE = Path(__file__).resolve().parent
ANKETA = BASE / 'consumer-questionnaire.xlsx'
SRC_DIR = BASE / 'scans-consumer'
OUT_DIR = BASE / 'output' / 'ocr-consumer'

DOC_EXT = fam.DOC_EXT  # pdf/jpg + Word (doc/docx/rtf/odt) + Excel (xlsx/xls/ods) + txt/csv

thin = Side(style='thin', color='999999')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FILL = PatternFill('solid', fgColor='DDEBF7')
TEST_FILL = PatternFill('solid', fgColor='FFF2CC')
HEAD_FONT = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical='top')

DELO_COLS = ['№', 'Суд (шапка)',
             'Истец ФИО', 'Истец дата рождения', 'Истец адрес регистрации',
             'Истец телефон', 'Паспорт истца (кем и когда выдан)',
             'Ответчик (наименование организации/ИП)', 'Адрес ответчика',
             'ИНН/ОГРН ответчика',
             'Стоимость товара (услуги), руб.', 'Дата покупки (заказа)',
             'Дата предъявления требования', 'Дата нарушения срока (отказа)',
             'Неустойка, руб.', 'Моральный вред, руб.', 'Прочие убытки, руб.',
             'Что просим (возврат денег/замена товара/устранение недостатков)',
             'Госпошлина, руб.', 'Приложения (свой список, через «;»)']

SKAN_COLS = ['Файл', 'Тип', 'Что извлечено / заметки', 'Дата']


def ocr_doc(path: Path) -> str:
    """Любой поддержанный документ: сканы — OCR, офисные — напрямую (fam.doc_text)."""
    return fam.doc_text(path, OUT_DIR)


def classify_doc(text, filename):
    f = filename.lower()
    if 'пасп' in f:
        return 'паспорт'
    if 'претенз' in f:
        return 'претензия'
    if 'чек' in f or 'касс' in f:
        return 'чек'
    t = text.upper()
    if 'ПАСПОРТ' in t or 'КЕМ ВЫДАН' in t:
        return 'паспорт'
    if 'ПРЕТЕНЗИЯ' in t or 'ТРЕБОВАНИЮ' in t:
        return 'претензия'
    if 'КАССОВЫЙ' in t or 'ЧЕК' in t:
        return 'чек'
    return ''


DEMO = {
    'Дело': ['1',
             'В Примерный городской суд\nПримерного района Примерной области',
             'Смирнова Ольга Петровна', '12.05.1987',
             'Примерная область, г. Пример, ул. Тестовая, д. 5, кв. 12',
             '8-911-000-00-00',
             '4010 123456 Отд. по вопросам миграции МВД России по Примерной области в Примерном районе от 20.02.2010',
             'ООО «ТехноМир»', 'Примерная обл., г. Пример, пр. Тестовый, д. 10', '7801234567891',
             '45990', '05.07.2025', '10.08.2025', '21.08.2025',
             '4599', '10000', '0',
             'возврат денег', '', 'копия товарного чека; копия претензии; фото недостатков'],
    'Проблематика': [
        '05.07.2025 в ООО «ТехноМир» мною приобретён смартфон марки «Тест-7» стоимостью '
        '45 990 рублей, что подтверждается товарным чеком. Гарантийный срок — 12 месяцев.',
        'В процессе эксплуатации, 02.08.2025, обнаружился недостаток: не работает приём '
        'Wi-Fi, телефон самопроизвольно перезагружается. Недостаток возник до передачи '
        'товара потребителю и не связан с нарушением правил эксплуатации.',
        '10.08.2025 в адрес продавца направлена письменная претензия с требованием '
        'возвратить уплаченную сумму (ст. 18, 22 Закона о защите прав потребителей).',
        'В установленный ст. 22 Закона о защите прав потребителей 10-дневный срок '
        'требование не удовлетворено: 21.08.2025 получен отказ, денежных средств до '
        'настоящего времени не поступило.',
    ],
}


def ensure_anketa(path: Path = ANKETA, demo: bool = False):
    """Создаёт анкету, если её нет. Существующую не трогает."""
    if path.exists():
        return False
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Инструкция'
    lines = [
        ['АНКЕТА ПОТРЕБИТЕЛЬСКОГО ДЕЛА — ИНСТРУКЦИЯ'],
        [''],
        ['Как работать:'],
        ['1. Сканы кладите в папку «scans-consumer»:'],
        ['   «паспорт истца» (распознаётся автоматически), «претензия», «чек», договор…'],
        ['   Тексты претензий и чеков сохраняются на лист «Сканы» — суммы и даты оттуда'],
        ['   впишите в лист «Дело» (генератор их не считает, он только сверяет).'],
        ['2. OCR вписывает данные паспорта истца ТОЛЬКО в пустые ячейки — ваши правки не сотрёт.'],
        ['3. Лист «Дело»: суд, ответчик (наименование ИП/ООО), стоимость товара, даты,'],
        ['   суммы неустойки/морального вреда, что просим.'],
        ['4. Лист «Проблематика» — ОПИСАТЕЛЬНАЯ ЧАСТЬ ИСКА, свободный текст:'],
        ['   каждая заполненная строка колонки A = отдельный абзац иска.'],
        ['   Опишите: что куплено, когда, где; какие недостатки, когда обнаружены;'],
        ['   когда и как предъявляли требование; что ответил продавец.'],
        ['5. В приложении — «Собрать исковое заявление». Проверьте «validation-report-consumer.txt».'],
        [''],
        ['Форматы: даты ДД.ММ.ГГГГ; суммы числами, без «рублей».'],
        ['Госпошлина: истец-потребитель освобождён по искам до 1 000 000 ₽'],
        ['(п. 3 ст. 333.36 НК РФ) — поле можно оставить пустым (в иске будет отметка об освобождении).'],
        [''],
        ['Жёлтые строки — ТЕСТОВЫЕ ВЫМЫШЛЕННЫЕ данные. Удалите их и вносите реальные.'],
    ]
    for row in lines:
        ws.append(row)
    ws.column_dimensions['A'].width = 110
    ws['A1'].font = Font(bold=True, size=14)

    ws = wb.create_sheet('Дело')
    ws.append(DELO_COLS)
    if demo:
        ws.append(DEMO['Дело'])
    for i, w in enumerate((5, 28, 24, 13, 30, 14, 34, 30, 30, 18,
                           15, 13, 14, 14, 13, 14, 13, 26, 12, 34), 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.font, cell.fill, cell.alignment, cell.border = HEAD_FONT, HEAD_FILL, WRAP, BORDER
    dv = DataValidation(type='list',
                        formula1='"возврат денег,замена товара,устранение недостатков"',
                        allow_blank=True)
    ws.add_data_validation(dv)
    dv.add('R2:R60')

    ws = wb.create_sheet('Проблематика')
    ws.append(['Описательная часть иска (свободный текст): каждая строка = абзац иска'])
    ws['A1'].font = HEAD_FONT
    ws.column_dimensions['A'].width = 120
    if demo:
        for t in DEMO['Проблематика']:
            ws.append([t])

    ws = wb.create_sheet('Сканы')
    ws.append(SKAN_COLS)
    for i, w in enumerate((36, 14, 60, 12), 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.font, cell.fill, cell.alignment, cell.border = HEAD_FONT, HEAD_FILL, WRAP, BORDER

    if demo:
        for row in wb['Дело'].iter_rows(min_row=2, max_row=2):
            for c in row:
                if c.value not in (None, ''):
                    c.fill = TEST_FILL
    wb.save(path)
    return True


def cell_empty(ws, row, col):
    v = ws.cell(row=row, column=col).value
    return v is None or str(v).strip() == ''


def fill_anketa(path: Path, results):
    """OCR-паспорт → только пустые ячейки листа «Дело»; остальные сканы — в журнал."""
    wb = openpyxl.load_workbook(path)
    ws, wj = wb['Дело'], wb['Сканы']
    heads = {str(c.value).strip(): c.column for c in ws[1] if c.value}
    changed = []

    def put_de(col_name, val):
        val = (val or '').strip()
        if not val:
            return
        col = heads.get(col_name)
        if col and cell_empty(ws, 2, col):
            ws.cell(row=2, column=col, value=val)
            changed.append(f'{col_name} ← {val}')

    doc = results.get('паспорт')
    if doc:
        put_de('Истец ФИО', doc.get('fio', ''))
        put_de('Истец дата рождения', doc.get('dr', ''))
        put_de('Истец адрес регистрации', doc.get('addr', ''))
        put_de('Паспорт истца (кем и когда выдан)',
               f'{doc.get("pasport", "")} {doc.get("kem_kogda", "")}'.strip(' ,.'))

    existing = {(str(wj.cell(row=r, column=1).value or '').strip(),
                 str(wj.cell(row=r, column=3).value or '').strip())
                for r in range(2, wj.max_row + 1)}
    for fname, tip, info in results.get('_rows', []):
        info = info[:4000]
        if (fname, info) in existing:
            continue
        r = 2 if wj.max_row == 1 else wj.max_row + 1
        wj.cell(row=r, column=1, value=fname)
        wj.cell(row=r, column=2, value=tip)
        wj.cell(row=r, column=3, value=info)
        existing.add((fname, info))
    wb.save(path)
    return changed


def scan_files(files):
    """OCR + классификация набора файлов → results для fill_anketa."""
    results = {'_rows': []}
    counts = {}
    total = len(files)
    print(f'[0/{total}] начало обработки', flush=True)
    for i, f in enumerate(files, 1):
        f = Path(f)
        try:
            text = ocr_doc(f)
        except Exception as ex:  # noqa: BLE001
            results['_rows'].append((f.name, '?', f'! не удалось прочитать: {ex}'))
            continue
        (OUT_DIR / platform_utils.ascii_filename(f.stem + '.txt')).write_text(text, encoding='utf-8')
        print(f'[{i}/{total}] {f.name}: готово', flush=True)
        tip = classify_doc(text, f.name)
        if tip == 'паспорт':
            p = fam.parse_passport(text)
            results['паспорт'] = p
            info = (f'ФИО: {p["fio"] or "?"}; {p["dr"] or "?"}; '
                    f'паспорт {p["pasport"] or "?"}; адрес: {p["addr"] or "?"}')
        elif tip == 'претензия':
            from datetime import datetime
            ds = fam.DATE_RE.findall(text)
            dates = [datetime.strptime(d.replace(',', '.'), '%d.%m.%Y') for d in ds]
            info = ('текст претензии сохранён ниже. Найденные даты: '
                    + ', '.join(d.replace(',', '.') for d in ds[:8]))
            if len(dates) >= 2:
                info += ' (обычно: дата покупки, дата претензии, срок ответа)'
            del dates
        elif tip == 'чек':
            sums = re.findall(r'(?:ИТОГО|итог|сумма|SUM)[:.\s-]*([\d ]+[,\.]\d{2})', text)
            info = ('текст чека сохранён ниже. Похоже на итог: '
                    + (sums[0] if sums else 'не найдено — впишите сумму вручную'))
        else:
            info = ('! тип не опознан — текст сохранён на лист «Сканы»; '
                    'для паспорта переименуйте файл: «паспорт истца»')
        results['_rows'].append((f.name, tip or '?', ' | '.join([info, text.strip()])[:6000]))
        counts[tip or '?'] = counts.get(tip or '?', 0) + 1
    return results, counts


def main():
    platform_utils.fix_console()
    not_found = platform_utils.missing_tools()
    if not_found:
        print('Для OCR не хватает программ:\n  • ' + '\n  • '.join(not_found))
        sys.exit(2)

    flags = {a for a in sys.argv[1:] if a.startswith('--')}
    demo = '--демо' in flags or '--demo' in flags
    demo_parser_run = '--демо-парсер' in flags
    args = [a for a in sys.argv[1:] if not a.startswith('--')]

    created = ensure_anketa(demo=demo)
    if created:
        print('Создана анкета: ' + ANKETA.name)
        if demo:
            print('  …с ТЕСТОВОЙ строкой (жёлтая) — удалите и вносите реальные данные.')
    else:
        print(f'Анкета на месте: {ANKETA.name}')

    if demo or demo_parser_run:
        if not args:
            print('Демо-режим: анкета готова, запустите «Собрать исковое заявление».')
            return

    SRC_DIR.mkdir(exist_ok=True)
    if args:
        files = []
        for a in args:
            p = Path(a) if Path(a).exists() else BASE / a
            files += (sorted(p.iterdir()) if p.is_dir() else [p])
        files = [p for p in files if p.exists()
                 and p.suffix.lower() in DOC_EXT and not p.name.startswith('~')]
    else:
        SRC_DIR.mkdir(exist_ok=True)
        files = sorted(p for p in SRC_DIR.iterdir()
                       if p.suffix.lower() in DOC_EXT and not p.name.startswith('~'))
    if not files:
        print('Документы (jpg/png, pdf, Word: doc/docx/rtf/odt, Excel: xlsx/xls/ods, '
          'txt/csv) не найдены в выбранной папке — проверьте и повторите.')
        return

    print(f'── РАСПОЗНАВАНИЕ ДОКУМЕНТОВ: {len(files)} файл(ов) '
          '(сканы — OCR, Word/Excel — напрямую) ──')
    results, counts = scan_files(files)
    changed = fill_anketa(ANKETA, results)
    print(f'Дозаполнено полей: {len(changed)} (только пустые ячейки)')
    for c in changed:
        print('  ' + c)
    for fname, tip, info in results.get('_rows', []):
        head = info.split(' | ')[0]
        print(f'  {fname} → {tip}: {head[:150]}')
    print('Сверьте распознанное с оригиналами — лист «Сканы» в анкете.')
    print('Данные не покидали этот компьютер.')


if __name__ == '__main__':
    main()
